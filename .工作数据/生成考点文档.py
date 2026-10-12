import json,re
from pathlib import Path
from collections import defaultdict
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_COLOR_INDEX
W=Path(__file__).resolve().parent;O=Path('/home/zyf/考研/四校复核新版');D=json.loads((W/'复核考点.json').read_text())
notes={
 '安徽大学':'基础卷2005与2007题目高度重复，本册将这两个年份的基础卷证据合为一个来源组，仍保留两年出处。读写卷单独统计。原始资料含填答与解析，填空题仅摘要列出考点，避免将资料中的填答当作标准答案。现有资料止于2021年；2017等卷有分组或任选题，出现次数不等于必答次数。',
 '黑龙江大学':'2011至2014分卷方式与后期不同，本册按现汉、古汉、语理知识部分分别汇总，同一学科同年最多计一次。2025、2026合并回忆资料已拆为学科证据。2020卷末未注明年份的现汉附段，以及2019古理末尾疑似现汉题不计。2022缺卷，2024现汉缺卷。',
 '云南民族大学':'现汉、语理分别计年次，共有知识点可在两科分别出现。2011、2013等早期语理及若干近年回忆资料不完整；近期数字仅表示可确认的出现，不能据未见题目推断未考。',
 '西安工业大学':'语综现有2020至2025年，评论现有2020至2026年。评论第一大题中古典与现代诗歌是任选关系，合记为一个鉴赏题型；第二、三大题也有材料选择。题型连续出现不等于某位作家或某首作品高频。'}
def font(r,size=10,bold=False,color='000000'):
 r.font.name='Noto Serif CJK SC';r.font.size=Pt(size);r.bold=bold;r.font.color.rgb=RGBColor.from_string(color)
 rf=r._element.get_or_add_rPr().find(qn('w:rFonts'))
 if rf is None:rf=OxmlElement('w:rFonts');r._element.get_or_add_rPr().insert(0,rf)
 for k in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+k),'Noto Serif CJK SC')
 return r
def p(d,text='',size=10,bold=False,keep=False,after=4,color='000000',style=None):
 x=d.add_paragraph(style=style);f=x.paragraph_format;f.line_spacing=1.15;f.space_before=Pt(0);f.space_after=Pt(after);f.keep_with_next=keep;f.widow_control=True;font(x.add_run(text),size,bold,color);return x
def setup(title):
 d=Document();s=d.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7);s.top_margin=Cm(1.7);s.bottom_margin=Cm(1.7);s.left_margin=Cm(1.8);s.right_margin=Cm(1.8);s.header_distance=Cm(.65);s.footer_distance=Cm(.7)
 for name in ['Normal','Title','Heading 1','Heading 2']:
  st=d.styles[name];st.font.name='Noto Serif CJK SC';st.font.color.rgb=RGBColor(0,0,0)
  for b in list(st.element.iter(qn('w:pBdr'))):b.getparent().remove(b)
 d.core_properties.title=title;d.core_properties.author='';d.core_properties.subject='2026年10月12日真题证据复核'
 h=s.header.paragraphs[0];h.alignment=WD_ALIGN_PARAGRAPH.RIGHT;font(h.add_run(title),8)
 f=s.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER;font(f.add_run('第 '),9);z=OxmlElement('w:fldSimple');z.set(qn('w:instr'),'PAGE');f._p.append(z);font(f.add_run(' 页'),9)
 p(d,title,19,True,after=10,style='Title');return d
def band(d,title):
 q=p(d,title,12,True,True,7,style='Heading 1');q.paragraph_format.space_before=Pt(8);q.alignment=WD_ALIGN_PARAGRAPH.CENTER
 z=OxmlElement('w:shd');z.set(qn('w:fill'),'F4E8BF');q._p.get_or_add_pPr().append(z)
def evidence_line(t):
 by=defaultdict(list)
 for e in t['evidence']:by[(e['year'],e['pages'])].append(e)
 return '；'.join(str(y)+' p'+pg for (y,pg) in sorted(by))
if __name__=='__main__':
 for school,ts in D.items():
  d=setup(school+'高频考点复核版')
  repeats=[t for t in ts if t['frequency']>=2];singles=[t for t in ts if t['frequency']==1]
  p(d,f'复现主题 {len(repeats)} 条  单次关注 {len(singles)} 条  2026年10月12日',9,color='555555')
  p(d,'先看复习顺序，再按章节练习。米色标章节，红色标题型，黄色为真题题目摘录或摘要。各条给出训练要求和全部已确认年份，便于回查。')
  p(d,'统计以同一学科的不同年份为单位，同年多题只计一次。次数为已核实证据的下限；“专题”汇总相关知识，“题型”统计任务形式，均不表示每个子点都重复。',9,color='555555')
  p(d,notes[school],9,color='555555')
  p(d,'“近年”取该科现有最新年份及此前四个自然年，仅列现存年份。缺卷与回忆遗漏会影响数字，故不换算命中率。p为本校《历年真题打印版》PDF文件页码；单页定位经题面文字匹配确认，未能唯一匹配的保留原卷页码范围。',9,color='555555')
  bydisc=defaultdict(list)
  for t in repeats:bydisc[t['discipline']].append(t)
  for k,tt in bydisc.items():
   pick=sorted([t for t in tt if t['kind']=='考点'],key=lambda t:(-len(t['recent_hits']),-t['frequency'],t['title']))[:5]
   if not pick:pick=sorted(tt,key=lambda t:-t['frequency'])[:3]
   p(d,k+'优先复习：'+'；'.join(t['title']+f"（{t['frequency']}年次）" for t in pick)+'。',9.5,after=6)
  sections=defaultdict(list)
  for t in repeats:sections[(t['discipline'],t['section'])].append(t)
  n=0
  for (k,sec),items in sorted(sections.items(),key=lambda v: (list(bydisc).index(v[0][0]), list(sections).index(v[0]))):
   band(d,sec if sec.startswith(('现代汉语','语言学理论','古代汉语','文学评论')) else k+' '+sec)
   for t in sorted(items,key=lambda t:(-len(t['recent_hits']),-t['frequency'],t['title'])):
    n+=1;q=p(d,'',keep=True,after=3,style='Heading 2');font(q.add_run(str(n)+'. '+t['title']),11,True)
    font(q.add_run('  '+t['kind']+'已确认 '+str(t['frequency'])+'年次'),9,True)
    q=p(d,'',keep=True,after=3);types='／'.join(sorted({e['type'] for e in t['evidence']}));font(q.add_run('【'+types+'】'),9.5,True,'A22F2F');font(q.add_run('掌握：'+t['points']),9.5)
    ev=t['evidence'];samples=[ev[0],ev[-1]] if len(ev)>1 else [ev[0]]
    q=p(d,'',keep=True,after=3)
    for i,e in enumerate(samples):
     if i:font(q.add_run('  '),9)
     tx=e['excerpt'][:66]+('…' if len(e['excerpt'])>66 else '')
     r=font(q.add_run(f"【{e['year']} {tx}】"),9,True);r.font.highlight_color=WD_COLOR_INDEX.YELLOW
    rys='、'.join(str(y) for y in t['recent_years']);hits='、'.join(str(y) for y in t['recent_hits']) or '暂无确认记录'
    p(d,'近年现存 '+rys+'；已见 '+hits+'。',8.5,color='555555',after=2,keep=True)
    lastp=p(d,'全部出处：'+evidence_line(t)+'。',8.5,color='555555',after=7);lastp.paragraph_format.keep_together=True
  band(d,'单次关注')
  p(d,'以下内容在本轮同科学段内仅确认一次，或经来源合并后只剩一组。保留它们用于覆盖复习，不列作高频。',9)
  for t in sorted(singles,key=lambda t:(t['discipline'],-max(t['years']),t['title'])):
   p(d,t['discipline']+'  '+t['title']+'  '+evidence_line(t),9.5,True,True,2)
   p(d,t['points'],9,after=5)
  p(d,'资料边界：统计基于目前可读题面、原图补证及人工归类，未对所有缺失题作补齐；页码与原题异常文字以打印复核版为准。',9,color='555555')
  path=O/(school+'_高频考点复核版.docx');d.save(path)
  assert n==len(repeats)
  text='\n'.join(x.text for x in Document(path).paragraphs)
  assert all(t['title'] in text for t in ts)
  print(path.name,'正文主题',n,'单次',len(singles),flush=True)
