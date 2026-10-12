import json,re
from pathlib import Path
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
W=Path(__file__).resolve().parent;O=W.parent/'四校难度评估';O.mkdir(exist_ok=True)
data=json.loads((W/'报告内容.json').read_text());loc=json.loads((W/'例题定位.json').read_text())
sources=[
 ('W1','云南民族大学2024初试参考书目 公开转载','https://ggw.100xuexi.com/WebSpecF/EnrolDetail.aspx?id=227519','核对719及815书目；815包含马学良语言学概论。'),
 ('W2','云南民族大学2026初试参考书目 公开转载','https://www.gaodun.com/kaoyan/1726539.html','核对新旧书目差别；不倒推旧卷范围。'),
 ('W3','黑龙江大学2015现代汉语大纲 公开转载','https://m.kaoyan.com/yanzhao/hlju/dagang/5443839dea9cd.html','旧书目为黄伯荣廖序东现代汉语，供历史卷参照。'),
 ('W4','黑龙江大学2025现代汉语大纲 公开转载','https://m.nutnow-lb.com/dagang/176013.html','近期书目转载列邵敬敏现代汉语通论，并列语义指向与歧义等范围；采用时保留转载身份。'),
 ('W5','黑龙江大学2026古代汉语与语言学理论大纲 公开转载','https://m.koolearn.com/kaoyan/20250930/1881800.html','核对王力一二册、修订版语言学纲要及古文阅读要求。'),
 ('W6','安徽大学2021文学院专业目录与范围 公开转载','https://jixun.iqihang.com/schoolzyml/20226650.html','与现有616及809卷相匹配；原官方历史链接本次无法读取。'),
 ('W7','西安工业大学2026招生考试参考书目 官方','https://grs.xatu.edu.cn/info/1010/4727.htm','确认816黄伯荣李炜现汉、郭锡良古汉及702童庆炳文学理论教程。'),
 ('W8','西安工业大学2024招生考试参考书目 官方','https://grs.xatu.edu.cn/info/1010/3627.htm','交叉核对902语言综合及702书目。')]

def font(r,size=10.5,bold=False,color='000000'):
 r.font.name='Noto Serif CJK SC';r.font.size=Pt(size);r.bold=bold;r.font.color.rgb=RGBColor.from_string(color)
 pr=r._element.get_or_add_rPr();rf=pr.find(qn('w:rFonts'))
 if rf is None:rf=OxmlElement('w:rFonts');pr.insert(0,rf)
 for k in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+k),'Noto Serif CJK SC')
 return r

def para(t='',size=10.5,bold=False,style=None,after=4,keep=False):
 p=d.add_paragraph(style=style);f=p.paragraph_format;f.line_spacing=1.12;f.space_before=Pt(0);f.space_after=Pt(after);f.keep_with_next=keep;f.widow_control=True
 font(p.add_run(t),size,bold);return p

def heading(t,level=1):
 p=para(t,14 if level==1 else 11.5,True,'Heading '+str(level),7,True);p.paragraph_format.space_before=Pt(7);return p

def table(headers,records,widths):
 t=d.add_table(rows=1,cols=len(headers));t.autofit=False
 for c,w in zip(t.columns,widths):c.width=Cm(w)
 for j,h in enumerate(headers):t.cell(0,j).text=h
 for row in records:
  cells=t.add_row().cells
  for j,tx in enumerate(row):cells[j].text=tx
 for ri,row in enumerate(t.rows):
  trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
  if ri==0:trpr.append(OxmlElement('w:tblHeader'))
  for j,c in enumerate(row.cells):
   c.width=Cm(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=c._tc.get_or_add_tcPr();margin=OxmlElement('w:tcMar')
   for side in ['top','left','bottom','right']:
    el=OxmlElement('w:'+side);el.set(qn('w:w'),'60');el.set(qn('w:type'),'dxa');margin.append(el)
   pr.append(margin)
   if ri==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DCE6F1');pr.append(sh)
   for p in c.paragraphs:
    f=p.paragraph_format;f.space_after=Pt(1);f.space_before=Pt(1);f.line_spacing=1.12;f.keep_with_next=False
    if j==len(headers)-1 and headers[-1]=='层级':p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs:font(r,9.5,ri==0)
 borders=OxmlElement('w:tblBorders')
 for k in ['top','left','bottom','right','insideH','insideV']:
  el=OxmlElement('w:'+k);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
 t._tbl.tblPr.append(borders)
 return t

def link(label,url):
 p=para('',9,after=3);h=OxmlElement('w:hyperlink');h.set(qn('r:id'),d.part.relate_to(url,RT.HYPERLINK,is_external=True));r=OxmlElement('w:r');pr=OxmlElement('w:rPr');rf=OxmlElement('w:rFonts');rf.set(qn('w:eastAsia'),'Noto Serif CJK SC');pr.append(rf);sz=OxmlElement('w:sz');sz.set(qn('w:val'),'18');pr.append(sz);r.append(pr);text=OxmlElement('w:t');text.text=label;r.append(text);h.append(r);p._p.append(h)

d=Document();s=d.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7);s.left_margin=Cm(1.8);s.right_margin=Cm(1.8);s.top_margin=Cm(1.6);s.bottom_margin=Cm(1.6);s.header_distance=Cm(.65);s.footer_distance=Cm(.65)
for name in ['Normal','Title','Heading 1','Heading 2']:
 st=d.styles[name];st.font.name='Noto Serif CJK SC';st.font.color.rgb=RGBColor(0,0,0)
 for el in list(st.element.iter(qn('w:pBdr'))):el.getparent().remove(el)
d.core_properties.title='四校真题相对参考书目难度评估';d.core_properties.subject='结果总览与逐校证据';d.core_properties.author=''
h=s.header.paragraphs[0];h.alignment=WD_ALIGN_PARAGRAPH.RIGHT;font(h.add_run('四校真题  相对参考范围评估'),8.5,color='555555')
p=s.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;font(p.add_run('第 '),9);field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');p._p.append(field);font(p.add_run(' 页'),9)
para('四校真题相对参考书目难度评估',20,True,'Title',10,True)
para('先看结果  再看标准与依据  2026年10月12日',9,after=8)
para(data['overview'],11,True,after=8)
para('本报告只评题目相对本校参考书目和考试范围的要求。比较基准是：已经读懂该校规定教材的主干知识，并完成典型基础练习后，还需要多少辨析、综合和陌生材料迁移。',10.5,after=8)
table(['学校','科目常见要求层级','相对范围的判断'],data['summary'],[2.9,5.1,9.4])
para('',after=2)
heading('四个最直接的判断',2)
for i,t in enumerate(data['final_points'],1):para(f'{i}. {t}',10,after=5)
heading('发散程度的比较',2)
para('开放论证与陌生文学文本迁移：西安的评论卷 ＞ 安徽的读写卷和部分文学论述 ＞ 黑龙江 ＞ 云南。教材内语言分析的精细度：黑龙江较突出。两种排序衡量的能力不同，不能用一个学校总分掩盖科目差异。',10)
para('这里的级别是命题要求的定性归类，不是统计出来的难度系数，也不代表分数预测。',9)

heading('评定标准与证据边界').paragraph_format.page_break_before=True
para('先确定知识点属于哪项规定范围，再看题目怎样要求运用它。知识点属于基础知识，仍可能被组合成较难的题；新句子、新作品也可能只是范围内的运用材料。')
table(['层级','判断标准'],data['criteria'],[4.6,12.8])
para('',after=2)
heading('本报告如何判断',2)
para('直接性：题目是否基本沿用教材概念和主干结论。分析深度：是否要区分相邻概念、说明条件并处理复杂语料。迁移程度：是否要自行选理论、选角度，并组织跨知识点或跨文本的论证。')
para('书目厚薄、科目门数、题量本身不直接加成难度。论述二字也不自动加成难度：一个教材原理的解释，可能只有2级；在陌生材料中选择论点并证明，才会上升到4级。')
heading('采用哪些真题',2)
para('安徽重点看2018至2021；黑龙江以2021、2023完整题面结合2025、2026回忆版；云南重点看2021至2025；西安语言综合看2020至2025，评论写作看2020至2026。早期资料用于校验常见问法，不机械拉平所有年代。')
para('例题页码指迁移包内各校《历年真题打印版》PDF的文件页码。材料中的题目摘录与摘要均来自现有文件，主表里的等级及下面的难易判断由本报告作出。未计算基础题或发散题的精确百分比。')
heading('书目核验的实际限制',2)
para('西安书目直接核对学校官网；安徽历史范围、云南相应书目、黑龙江大纲正文采用公开转载，并保留身份。黑龙江官网确认发布了2026大纲，但附件下载要求验证码，本次未取得附件原文。安徽历史官方原链接本次无法读取。')
para('没有取得全部教材全文，也未逐题逐页验证书上原题。因此，本报告可以评价教材主题的直接提取、方法运用和迁移要求，不能宣称已经证明所有题目都在某一页或没有任何超纲题。2027新书目不倒套到现有历史卷。',9.5)

for school in data['schools']:
 heading(school['name']).paragraph_format.page_break_before=True
 para(school['verdict'],11.5,True,after=7)
 para('参考范围  '+school['scope'],9.5,after=6)
 para(school['claim'],10,after=6)
 heading('代表题目与判断依据',2)
 records=[]
 for pid,exam,q,why,grade in school['cases']:
  e=loc[school['name']+':'+str(pid)];pages=e['pages']
  if school['name']=='西安工业大学' and pid==730:pages=[55,56]
  if school['name']=='西安工业大学' and pid==773:pages=[56,57]
  if school['name']=='西安工业大学' and pid==774:pages=[57,58,59,60]
  pg=str(pages[0]) if len(pages)==1 else str(min(pages))+'至'+str(max(pages))
  records.append([exam+'\n'+q+'\n原题第'+pg+'页',why,grade])
 table(['真题与出处','对应范围和能力要求','层级'],records,[6.1,9.4,1.9])
 para('',after=2)
 heading('为什么这样评',2);para(school['analysis'],10,after=5)
 para(school['boundary'],9.5,after=5)
 para(school['conclusion'],10.5,True,after=4)

heading('如何理解这些结论').paragraph_format.page_break_before=True
para('四校主干都建立在本科教材知识和相应训练之上。区别主要在于：云南让你把主干知识说清、用出来；黑龙江让你把边界和理由分清；安徽让你把概念与作品及材料连成解释；西安评论卷让你在给定作品中选择理论并形成文章。',10.5)
para('如果把熟读参考书理解为记住概念，那么黑龙江的分析题、安徽的综合题和西安的评论题仍会有明显障碍。如果熟读还包括理解例证、完成课后分析及用理论解释文本，则这些要求多数仍可由规定学习范围提供基础。',10.5)
para('不能把书里没有这道原题等同于超出书目范围。真正需要另行核对的，是题目是否要求使用书目与大纲均未覆盖的概念、方法或必须预先掌握的专门材料。现有证据支持按命题方式区分难度，不支持给四校编造超纲比例。',10.5)
heading('书目与大纲来源',2)
for code,title,url,note in sources:
 link('['+code+'] '+title,url);para(note,8.5,after=3)
heading('本地真题来源',2)
para('四校真题换电脑交接／四校真题打印复核版中的安徽大学、黑龙江大学、云南民族大学、西安工业大学四份PDF。对应题面已在前文逐条给出年份、科目及文件页码。来源检查日期为2026年10月12日。',9)
para('统计与判断的主要限制：安徽缺2022年以后的对应旧科目卷；黑龙江近两年为回忆版；云南含提要或缺语料；西安缺2026语言综合卷。此报告评现有卷的命题要求，不能据此断定未见年份的变化。',9)

path=O/'四校真题相对参考书目难度评估.docx';d.save(path)
texts='\n'.join(p.text for p in Document(path).paragraphs)
assert data['overview'] in texts and all(s['conclusion'] in texts for s in data['schools'])
assert len(d.tables)==6 and sum(len(s['cases']) for s in data['schools'])==22
assert all((W.parent/'四校真题换电脑交接'/'四校真题打印复核版'/(s['name']+'_历年真题_打印版.pdf')).is_file() for s in data['schools'])
print(path,'22条定位例题，四校结论及来源全部写入。')
