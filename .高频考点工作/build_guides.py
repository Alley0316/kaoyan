import json,re
from pathlib import Path
from collections import defaultdict
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_BREAK,WD_COLOR_INDEX

W=Path(__file__).resolve().parent
O=W.parent/'四校高频考点';O.mkdir(exist_ok=True)
data=json.loads((W/'curated_topics.json').read_text())
index=json.loads((W/'pdf_index.json').read_text())
short={'安徽大学':'安大','黑龙江大学':'黑大','云南民族大学':'云民大','西安工业大学':'西安工大'}
course={'语言文学基础':'基础','阅读与写作':'读写','语言学理论':'语理','现代汉语':'现汉','语言学理论与现代汉语语音文字词汇':'理现','古代汉语与现代汉语语法修辞':'古现','古代汉语与语言学理论':'古理','古代汉语与语言学理论 回忆版':'古理','现代汉语和古代汉语与语言学理论 回忆版':'合卷','语言综合':'语综','文学评论写作':'评论'}
label={'名词':'名词解释','问答':'简答问答','论述':'论述','分析':'分析','阅读':'阅读翻译','鉴赏':'鉴赏','写作':'写作','运用':'运用','评论':'评论','填空':'填空','题目':'综合','未明':'题型未明'}
notes={
 '安徽大学':'安徽2005、2007基础卷在现有资料中大量重复，本册按资料标注分别计卷，不能据此确认两年的独立出题情况。2012等卷含回忆性表述；原始资料附有答案，本册仅用题目作为复现依据。',
 '黑龙江大学':'2024至2026含手写或文字回忆版，未回忆出的题目不补入。2020古理卷末附有未单独标注年份的现汉题段，本册未将该段计入频次。现有资料缺2022卷；部分古文有模糊或字行重叠。',
 '云南民族大学':'现有资料的年份、科目不连续，若干卷为回忆或题目提要，部分语料缺失。2019语理前五题未标题型，以“未”标注。频次只反映已能确认的题目；分析题缺例词时仍可确认题型，不能补造例词。',
 '西安工业大学':'语言综合现有2020至2025卷，文学评论写作现有2020至2026卷；未收录2026语言综合。评论卷中部分理论材料以原图保留，本册只统计可确认的题型及材料。'}

def font(run,size=11,bold=False,color='000000'):
 run.font.name='Noto Serif CJK SC';run.font.size=Pt(size);run.bold=bold;run.font.color.rgb=RGBColor.from_string(color)
 rpr=run._element.get_or_add_rPr();rf=rpr.find(qn('w:rFonts'))
 if rf is None:rf=OxmlElement('w:rFonts');rpr.insert(0,rf)
 for k in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+k),'Noto Serif CJK SC')
 for k in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:rf.attrib.pop(qn('w:'+k),None)
 return run

def pf(p,after=3,keep=False):
 f=p.paragraph_format;f.line_spacing=1.1 if school in ('黑龙江大学','西安工业大学') else 1.15;f.space_before=Pt(0);f.space_after=Pt(after);f.keep_with_next=keep;f.widow_control=True

def plain(d,t,size=10,color='000000',after=4,keep=False):
 p=d.add_paragraph();pf(p,after,keep);font(p.add_run(t),size,color=color);return p

for school,topics in data.items():
 d=Document();s=d.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7);s.top_margin=Cm(1.7);s.bottom_margin=Cm(1.7);s.left_margin=Cm(1.8);s.right_margin=Cm(1.8);s.header_distance=Cm(.7);s.footer_distance=Cm(.7)
 for name in ['Normal','Title','Heading 1','Heading 2']:
  st=d.styles[name];st.font.name='Noto Serif CJK SC';st.font.color.rgb=RGBColor(0,0,0);st.font.size=Pt(11)
  st.element.get_or_add_rPr().set(qn('w:lang'),'zh-CN')
  for border in list(st.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
 d.core_properties.title=school+'高频考点';d.core_properties.subject='按现有历年真题归并的复习提纲';d.core_properties.author='';d.core_properties.keywords='历年真题 高频考点 复习提纲'
 h=s.header.paragraphs[0];h.alignment=WD_ALIGN_PARAGRAPH.RIGHT;font(h.add_run(school+'  高频考点'),8,color='555555')
 foot=s.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER;font(foot.add_run('第 '),9)
 field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field);font(foot.add_run(' 页'),9)
 p=d.add_paragraph(style='Title');pf(p,8,True);font(p.add_run(school+'高频考点'),19,True)
 ee=index[school];plain(d,f"现有{ee[0]['year']}至{ee[-1]['year']}年资料  {len(ee)}份卷次  {len(topics)}条复现主题  整理日期 2026年10月12日",9,'555555',6)
 plain(d,'按章节复习，每条先掌握概念与辨析，再练习对应题型。黄色标出真题摘录或题目摘要；出处中的页码均指本校《历年真题打印版》PDF的文件页码。',10)
 plain(d,'统计口径：按标注年份和科目计不同卷次，同卷同主题多题只计1卷。5卷及以上标为高频，3至4卷标为多次，2卷标为复现。宽主题合并相关子点，频次不表示每个子点都被重复考过。',9,'555555')
 plain(d,notes[school],9,'555555')
 legends={course[x['subject']]:x['subject'].replace(' 回忆版','') for x in ee}
 plain(d,'科目简称：'+'；'.join(k+'＝'+v for k,v in legends.items())+'。名＝名词解释；问＝简答问答；析＝分析；论＝论述；读＝阅读翻译；鉴＝鉴赏；写＝写作；运＝运用；评＝评论；填＝填空；未＝题型未明。',8.5,'555555',7)
 priority=sorted(topics,key=lambda x:-x['frequency'])[:5]
 plain(d,'优先复习：'+'；'.join(x['title']+f"（{x['frequency']}卷）" for x in priority)+'。',9.5,after=8)
 sections=defaultdict(list)
 for t in topics:sections[t['section']].append(t)
 n=0
 for section,items in sections.items():
  p=d.add_paragraph(style='Heading 1');pf(p,7,True);p.paragraph_format.space_before=Pt(8);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;font(p.add_run(section),13,True)
  shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'F4E8BF');p._p.get_or_add_pPr().append(shade)
  for t in sorted(items,key=lambda x:-x['frequency']):
   n+=1;rank='高频' if t['frequency']>=5 else ('多次' if t['frequency']>=3 else '复现')
   p=d.add_paragraph(style='Heading 2');pf(p,3,True);p.paragraph_format.space_before=Pt(5)
   font(p.add_run(str(n)+'. '),11,True)
   font(p.add_run('【'+'／'.join(label[x] for x in t['types'])+'】'),10,True,'A22F2F')
   font(p.add_run(t['title']),11,True)
   font(p.add_run(f"  {rank} {t['frequency']}卷 {t['years']}年"),9.2,True)
   plain(d,'掌握：'+t['points'],9.8,after=3,keep=True)
   ev=t['evidence'];first=min(ev,key=lambda x:(x['year'],x['eid'],x['pid']));last=max(ev,key=lambda x:(x['year'],x['eid'],x['pid']))
   # Use specific titles to distinguish the repeated classical and modern poetry tasks.
   if school=='西安工业大学' and t['title']=='中国古典诗词鉴赏':
    first=first.copy();first['excerpt']='李商隐 无题／李白 行路难 任选其一鉴赏';last=last.copy();last['excerpt']='李商隐 锦瑟／欧阳修 蝶恋花 任选其一鉴赏'
   if school=='西安工业大学' and t['title']=='中国现代诗歌鉴赏':
    first=first.copy();first['excerpt']='余光中 乡愁 鉴赏';last=last.copy();last['excerpt']='戴望舒 夕阳下 鉴赏'
   if school=='西安工业大学' and t['title']=='文论材料阐释':
    first=first.copy();first['excerpt']='王国维 一代有一代之文学 理论阐释';last=last.copy();last['excerpt']='兴观群怨／自己的文章／新的美学原则在崛起 材料阐释'
   if school=='西安工业大学' and t['title']=='外国文学评论':
    first=first.copy();first['excerpt']='法的门前／我为何而生 任选材料写评论';last=last.copy();last['excerpt']='荒原／百年孤独 任选材料写评论'
   if school=='安徽大学' and t['title']=='诗歌鉴赏写作':
    first=first.copy();first['excerpt']='江雪／鹧鸪天 任选一题赏析';last=last.copy();last['excerpt']='小溪 赏析'
   p=d.add_paragraph();pf(p,3,True)
   for i,e in enumerate([first,last]):
    e=e.copy()
    summaries={203:'填写辅音的国际音标并描写元音',285:'双音节语素的类型',58:'列出三十六字母',233:'举例说明假借与通假的区别',234:'古代汉语被动标志的组合',16:'单纯词与合成词的区别'}
    if e['pid'] in summaries and school==('安徽大学' if e['pid'] in (203,285) else ('西安工业大学' if e['pid']==16 else '黑龙江大学')):e['excerpt']=summaries[e['pid']]
    e['excerpt']=re.sub(r'(?:郭锡良|黄廖本|教辅|教材)[上下本]?[Pp]\d+(?:-\d+)?','',e['excerpt'])
    if i:font(p.add_run('  '),9.5)
    text=f"【{short[school]}{e['year']} {course[e['subject']]} {e['type']} {e['excerpt']}】"
    run=font(p.add_run(text),9.5,True);run.font.highlight_color=WD_COLOR_INDEX.YELLOW
   byexam=defaultdict(list)
   for e in ev:byexam[e['eid']].append(e)
   typename_short={'名词':'名','问答':'问','分析':'析','论述':'论','阅读':'读','鉴赏':'鉴','写作':'写','运用':'运','评论':'评','填空':'填','题目':'综','未明':'未'}
   pieces=[]
   for eid,es in sorted(byexam.items(),key=lambda x:(x[1][0]['year'],x[0])):
    e=es[0];tp='/'.join(typename_short[a] for a in sorted({x['type'] for x in es}));pieces.append(f"{str(e['year'])[2:]}{course[e['subject']]}{tp}p{e['pages']}")
   p=plain(d,'历年出处：'+'；'.join(pieces)+'。',8.5,'555555',3 if school in ('黑龙江大学','西安工业大学') else 6)
   p.paragraph_format.keep_together=True
 # Single coherent locator to the unchanged source document.
 plain(d,'核对原题：迁移包内 四校真题打印复核版／'+school+'_历年真题_打印版.pdf。黄色内容供定位和复习；需完整题干、原图或原件异常用词时，以该文件对应页为准。',9,'555555',6)
 plain(d,'这些频次反映现有资料中的重复考查，不能推定缺失年份，也不能视为下一次考试的命中承诺。',9,'555555')
 path=O/(school+'_高频考点.docx');d.save(path)
 # Small non-trivial builder checks: all topics and evidence counts remain in the editable artifact.
 text='\n'.join(p.text for p in Document(path).paragraphs)
 assert all(x['title'] in text and f"{x['frequency']}卷 {x['years']}年" in text for x in topics)
 assert n==len(topics)
 print(path.name,len(topics),'条',len(text),'字',flush=True)
