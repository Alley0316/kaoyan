import json,re,hashlib,shutil
from pathlib import Path
from docx import Document
import pdfplumber
W=Path(__file__).resolve().parent;O=W.parent/'四校高频考点';topics=json.loads((W/'curated_topics.json').read_text())
def norm(s):return re.sub(r'\s+','',s)
report={}
for school,tt in topics.items():
 source=W/'render'/school/(school+'_高频考点.pdf');target=O/source.name
 shutil.copyfile(source,target)
 with pdfplumber.open(target) as pdf:
  text='\n'.join(p.extract_text() or '' for p in pdf.pages);nt=norm(text)
  doc=Document(O/(school+'_高频考点.docx'))
  missing=[p.text for p in doc.paragraphs if p.text and norm(p.text) not in nt]
  assert not missing,(school,'Word文字未出现在PDF',missing)
  assert sum(p.text.startswith('掌握：') for p in doc.paragraphs)==len(tt)
  assert text.count('掌握：')==len(tt)
  for t in tt:
   assert t['frequency']==len({e['eid'] for e in t['evidence']})
   assert t['years']==len({e['year'] for e in t['evidence']})
  for n,p in enumerate(pdf.pages,1):
   assert abs(p.width-595.28)<1 and abs(p.height-841.89)<1
   assert (W/'render'/school/f'page-{n}.png').is_file()
   body=[c for c in p.chars if 45<c['top']<p.height-40]
   assert min(c['x0'] for c in body)>=49 and max(c['x1'] for c in body)<=p.width-49,(school,n,'正文边界')
   assert '\ufffd' not in (p.extract_text() or '') and '\u25a0' not in (p.extract_text() or ''),(school,n,'替换字符')
  # Remove only obsolete render images from earlier pagination, never sources or deliverables.
  for p in (W/'render'/school).glob('page-*.png'):
   if int(p.stem.split('-')[-1])>len(pdf.pages):p.unlink()
  report[school]={'topics':len(tt),'pages':len(pdf.pages),'docx_text_preserved_in_pdf':True,'a4_bounds_check':True,'sha256':{p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in [O/(school+'_高频考点.docx'),target]}}
print(json.dumps({s:{'考点':v['topics'],'页数':v['pages']} for s,v in report.items()},ensure_ascii=False))
(W/'最终检查.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
print('194条主题及全部Word正文已保留；28页A4尺寸、正文边界、异常字符检查通过。')
